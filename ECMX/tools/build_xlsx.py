#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ECMX_materials.xlsx 생성 (§9-5).
시트: Essential / Tissue_Matrix / Inventory / By_Layer / Suppliers / Patents / Evidence / Confidence / Gaps
"""
import csv, os, re, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/mnt/user-data/outputs/ECMX"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
TIER_FILL = {"E0": "C00000", "E1": "ED7D31", "E2": "FFD966", "E3": "D9D9D9"}
TIER_FONT = {"E0": "FFFFFF", "E1": "FFFFFF", "E2": "000000", "E3": "000000"}
CONF_FILL = {"상": "C6EFCE", "중": "FFEB9C", "하": "FFC7CE"}
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def read(fn):
    p = os.path.join(BASE, fn)
    if not os.path.exists(p):
        return [], []
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return (list(rows[0].keys()) if rows else []), rows


def sheet(wb, name, cols, rows, widths=None, note=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, size=9, color="666666")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(cols)))
        r0 = 2
    for j, c in enumerate(cols, 1):
        cell = ws.cell(r0, j, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    for i, row in enumerate(rows, r0 + 1):
        for j, c in enumerate(cols, 1):
            v = (row.get(c) or "") if isinstance(row, dict) else (row[j - 1] if j - 1 < len(row) else "")
            cell = ws.cell(i, j, v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(size=9); cell.border = THIN
            s = str(v).strip()
            if c in ("tier",) and s[:2] in TIER_FILL:
                cell.fill = PatternFill("solid", fgColor=TIER_FILL[s[:2]])
                cell.font = Font(size=9, bold=True, color=TIER_FONT[s[:2]])
            if c in ("confidence",) and s in CONF_FILL:
                cell.fill = PatternFill("solid", fgColor=CONF_FILL[s])
                cell.font = Font(size=9, bold=True)
    for j, c in enumerate(cols, 1):
        w = (widths or {}).get(c, 0)
        if not w:
            mx = max([len(str(c))] + [len(str((r.get(c) or "") if isinstance(r, dict) else "")) for r in rows[:250]] or [10])
            w = min(52, max(11, int(mx * 0.85)))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(max(1,len(cols)))}{r0+len(rows)}"
    return ws


def main():
    wb = Workbook(); wb.remove(wb.active)

    # --- Essential (M3) ---
    c, r = read("essential_items.csv")
    sheet(wb, "Essential", c, r,
          {"experiment_summary": 56, "quant_range": 34, "evidence_ref": 28,
           "confidence_reason": 40, "alternatives": 30, "open_issues": 32, "name_ko": 20},
          "M3 오가노이드 지지체 필수 물질 — 등급 E0(절대필수)/E1(조건부)/E2(선택)/E3(대체가능·근거없음). "
          "E0·E1은 OMIT(제거)·SUBST(치환)·DOSE(농도반응) 실험 근거를 반드시 보유. 기준일 2026-09-14")
    ess = r

    # --- Tissue_Matrix ---
    c2, r2 = read("essential_matrix.csv")
    ws = sheet(wb, "Tissue_Matrix", c2, r2, {c2[0] if c2 else "조직": 22},
               "조직 × 요구사항 매트릭스. 셀 값 E0~E3. 색상만으로 구분하지 않도록 글자를 함께 표기했다.")
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            s = str(cell.value or "").strip()[:2]
            if s in TIER_FILL:
                cell.fill = PatternFill("solid", fgColor=TIER_FILL[s])
                cell.font = Font(size=9, bold=True, color=TIER_FONT[s])
    for j in range(2, len(c2) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 26

    # --- Inventory (M2 통합) ---
    inv_cols, inv = [], []
    for fn in ("inv_bio.csv", "inv_chem.csv", "inv_spec.csv"):
        cc, rr = read(fn)
        if cc and not inv_cols: inv_cols = cc
        inv += rr
    if inv_cols:
        sheet(wb, "Inventory", inv_cols, inv,
              {"function_role": 40, "spec": 34, "confidence_reason": 36, "risk": 30,
               "alternatives": 28, "synonyms": 24, "ecm_counterpart": 24, "price": 26},
              "M2 ECM 소재 개발용 물질 인벤토리 L1~L9 통합")
        # 통합본 저장
        with open(os.path.join(BASE, "material_inventory.csv"), "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=inv_cols); w.writeheader()
            for row in inv: w.writerow({k: row.get(k, "") for k in inv_cols})

        # --- By_Layer 피벗 ---
        LN = {"L1": "구조 단백질", "L2": "GAG·다당", "L3": "프로테오글리칸·연결분자",
              "L4": "성장인자·신호분자", "L5": "펩타이드 모티프", "L6": "가교제·개시제",
              "L7": "합성 백본·하이브리드", "L8": "완충·이온·부형제", "L9": "품질·안전 규격"}
        cnt = collections.Counter((x.get("layer") or "").strip() for x in inv)
        ess_cnt = collections.Counter(
            (x.get("layer") or "").strip() for x in inv if (x.get("essentiality") or "").strip() == "필수")
        conf = collections.Counter(
            ((x.get("layer") or "").strip(), (x.get("confidence") or "").strip()) for x in inv)
        rows = [[l, LN[l], cnt.get(l, 0), ess_cnt.get(l, 0),
                 conf.get((l, "상"), 0), conf.get((l, "중"), 0), conf.get((l, "하"), 0)] for l in LN]
        rows.append(["합계", "", sum(cnt.values()), sum(ess_cnt.values()),
                     sum(v for (l, g), v in conf.items() if g == "상"),
                     sum(v for (l, g), v in conf.items() if g == "중"),
                     sum(v for (l, g), v in conf.items() if g == "하")])
        sheet(wb, "By_Layer", ["층위", "층위명", "물질 수", "필수 등재", "신뢰도 상", "신뢰도 중", "신뢰도 하"],
              rows, {"층위명": 26}, "L1~L9 층위별 집계")

    # --- Suppliers / Patents ---
    c, r = read("suppliers.csv")
    if c: sheet(wb, "Suppliers", c, r, {"spec": 34, "supply_risk": 30, "source_url": 40,
                                        "confidence_reason": 32, "product_name": 30},
                "공급·조달. 가격은 공개 카탈로그 기준이며 대량 구매 조건과 다를 수 있다. 확인일 2026-09-14")
    pc, pr = read("patents_landscape.csv")
    pc2, pr2 = read("patent_essentials.csv")
    if pc: sheet(wb, "Patents", pc, pr, {"title": 40, "significance": 36, "expiry_note": 30,
                                         "confidence_reason": 30}, "특허 랜드스케이프 (에이전트 G)")
    if pc2: sheet(wb, "Patents_Essential", pc2, pr2,
                  {"title": 34, "essential_elements": 40, "composition_ranges": 38,
                   "claim_quote": 54, "confidence_reason": 28},
                  "특허 독립항 필수 구성요소 (에이전트 P). claim_quote는 청구항 원문 직접 인용")

    # --- Evidence ---
    lp = os.path.join(BASE, "evidence_ledger.md")
    ev = []
    if os.path.exists(lp):
        for line in open(lp, encoding="utf-8"):
            if line.startswith("|") and not re.match(r"^\|\s*[-:]+", line) and "제목" not in line:
                cells = [x.strip() for x in line.strip().strip("|").split("|")]
                if len(cells) >= 6: ev.append(cells)
    if ev:
        n = max(len(x) for x in ev)
        ev = [x + [""] * (n - len(x)) for x in ev]
        cols = (["#", "제목", "저널", "연도", "DOI", "PMID", "전문", "인용 파일"] if n == 8
                else [f"col{i+1}" for i in range(n)])
        sheet(wb, "Evidence", cols[:n], ev, {"제목": 60, "인용 파일": 30},
              "근거 대장 — 논문·특허 전체 (evidence_ledger.md 추출)")

    # --- Confidence ---
    allrows = []
    for src, rows_ in (("Essential", ess), ("Inventory", inv), ("Suppliers", r if c else []),
                       ("Patents", pr if pc else []), ("Patents_Essential", pr2 if pc2 else [])):
        for x in rows_:
            idk = x.get("essential_id") or x.get("material_id") or x.get("supplier_id") or x.get("patent_id") or ""
            nm = x.get("name_ko") or x.get("product_name") or x.get("title") or ""
            allrows.append([src, idk, nm[:70], (x.get("confidence") or "").strip(),
                            (x.get("confidence_reason") or "")[:230]])
    if allrows:
        ws = sheet(wb, "Confidence", ["산출물", "ID", "항목", "신뢰도", "판정 사유"], allrows,
                   {"항목": 44, "판정 사유": 62},
                   "항목 단위 신뢰도 등급과 판정 사유 (R-08). 상=1차출처 2건 이상 교차확인, 중=1차출처 1건, 하=2차출처·추정")
        for row in ws.iter_rows(min_row=3, min_col=4, max_col=4):
            for cell in row:
                if str(cell.value).strip() in CONF_FILL:
                    cell.fill = PatternFill("solid", fgColor=CONF_FILL[str(cell.value).strip()])
                    cell.font = Font(size=9, bold=True)

    # --- Gaps ---
    gp = os.path.join(BASE, "gaps.md")
    gaps = []
    if os.path.exists(gp):
        for line in open(gp, encoding="utf-8"):
            if line.startswith("|") and not re.match(r"^\|\s*[-:]+", line):
                cells = [x.strip() for x in line.strip().strip("|").split("|")]
                if len(cells) >= 3 and "항목" not in cells[0]: gaps.append(cells[:4])
    if gaps:
        n = max(len(x) for x in gaps)
        gaps = [x + [""] * (n - len(x)) for x in gaps]
        sheet(wb, "Gaps", (["구분", "항목", "실패 사유", "영향"])[:n], gaps,
              {"항목": 40, "실패 사유": 56, "영향": 34}, "확보 실패·결측 항목 (유예 표현 없이 실패로 기록)")

    out = os.path.join(BASE, "ECMX_materials.xlsx")
    wb.save(out)
    print("saved:", out, "| sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
